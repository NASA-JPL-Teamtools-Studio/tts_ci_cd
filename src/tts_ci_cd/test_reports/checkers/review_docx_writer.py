import pdb
from datetime import datetime
from docxtpl import DocxTemplate, RichText
from docx import Document
import re
from pathlib import Path
import logging

from tts_tower.checkers.checker_base import CheckerBase
from tts_tower.rule_results import RULE_REVISION_SPECIAL
from tts_html_utils.core.components.text import Paragraph as P
from tts_html_utils.core.components.text import Heading2 as H2
from tts_html_utils.core.components.text import Heading3 as H3

logger = logging.getLogger(f'tower.checkers.review_docx_writer')

ACCEPTABLE_ACVR_HARD_VIOLATIONS = ['OCO-AC-010']
ACCEPTABLE_ACVR_SOFT_VIOLATIONS = ['OCO-AC-008']

def extract_variables(docx_file):
    ###############################################################
    # Special note: this function is pretty obtuse. It came largely
    # from ChatGPT. But I think that's OK because it's only being
    # used to provide warnings. If anything it will come up with
    # false positives if/when it fails, so it will be wrapped
    # in a warning in the logs.
    ###############################################################
    doc = Document(docx_file)
    non_dict_variables = []  # To store non-dictionary variables
    dict_variables = {}  # To store dictionary-like variables

    ACCEPTABLE_ACVR_HARD_VIOLATIONS = ['OCO-AC-010']
    ACCEPTABLE_ACVR_SOFT_VIOLATIONS = ['OCO-AC-008']

    # Regex pattern to capture the iteration variable and collection variable in the for loop
    for_pattern = r'{% for\s+(\w+)\s+in\s+(\w+) %}'

    # Regex pattern to capture {{r variable}} or {{variable}} placeholders
    placeholder_pattern = r'{{\s*r?\s*(\w+)\s*}}'

    # Regex pattern to capture dictionary keys in {{r removed_target['Key']}}
    key_pattern = r"r?\s*(\w+)\s*\[\s*(.*?)\s*\]"

    last_collection_var = None

    def find_matches_in_text(text, last_collection_var):
        # Capture the loop variable and collection variable in the for loop (like removed_target in removed_targets)
        for_matches = re.findall(for_pattern, text)
        # If a loop variable and collection variable are found, initialize it in the dictionary
        for loop_var, collection_var in for_matches:
            last_collection_var = collection_var
            if collection_var not in dict_variables:
                dict_variables[collection_var] = []

        # Capture {{r variable}} and {{ variable }} placeholders (non-dict variables)
        placeholder_matches = re.findall(placeholder_pattern, text)
        for match in placeholder_matches:
            if match not in dict_variables:  # If not a dict-like variable
                non_dict_variables.append(match)
            else:
                dict_variables[match].append(match)

        # Capture dictionary keys in {{r removed_target['Key']}}
        key_matches = re.findall(key_pattern, text)
        for match in key_matches:
            variable_name, key = match
            if last_collection_var is None:
                raise Exception('How did I get here?')
            # If this key corresponds to a valid collection variable, add it to the dictionary
            if last_collection_var in dict_variables and key not in dict_variables[last_collection_var]:
                dict_variables[last_collection_var].append(key[1:-1]) #[1:-1] strips quotes

        return last_collection_var

    # Check all paragraphs (including headings) in the main document
    for para in doc.paragraphs:
        last_collection_var = find_matches_in_text(para.text, last_collection_var)

    # Check all tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                last_collection_var = find_matches_in_text(cell.text, last_collection_var)

    # Check headers and footers
    for section in doc.sections:
        # Check header
        header = section.header
        for para in header.paragraphs:
            last_collection_var = find_matches_in_text(para.text, last_collection_var)

        # Check footer
        footer = section.footer
        for para in footer.paragraphs:
            last_collection_var = find_matches_in_text(para.text, last_collection_var)

    return non_dict_variables, dict_variables



# ===================================================================
# :: Checker Class
# ===================================================================
class WriteAtsReport(CheckerBase):
    FR_IDS = [
        ('PIPE-000', RULE_REVISION_SPECIAL.USER_RULE, 'Ceci n\'est pas une pipe'),
    ]

    INPUTS = [
        ('acvr', True),
        ('ccvr', True),
        ('odf', True),
        ('nephthys', True),
        ('week_directory', True),
        ('scr', True),
        ('switch_buffer', True),
        ('abs', True),
        ('expanded_ats', True),
        ('mpst_team', True),
        ('violations_table', True),
        ('car', True),
        ('mnvr', True),
        ('oef', True)
    ]

    def _impl_init(self, acvr_client, ccvr_client, odf_client, nephthys_client, week_directory_client, scr_client, switch_buffer_client, abs_client, expanded_ats_client, mpst_team_client, violations_table, car_client, maneuver_client, oef_client):
        self.template_filepath = Path(__file__).parent.parent.joinpath('support_files/Review Notes ATS MPST Review Comments TEMPLATE.docx')
        self.output_filepath = f'{week_directory_client.output_path}/Review Comments for {scr_client.filepath.stem}.docx'
        self.context = {}

    def _impl_do_rulecheck(self, acvr_client, ccvr_client, odf_client, nephthys_client, week_directory_client, scr_client, switch_buffer_client, abs_client, expanded_ats_client, mpst_team_client, violations_table, car_client, maneuver_client, oef_client):
        template = DocxTemplate(self.template_filepath)
        expected_non_record_variables, expected_record_variables = extract_variables(self.template_filepath)

        #Sometimes we want to suppress warnings below. This gives us a hook to do it with
        excused_context_keys = {}
        excused_context_keys['tertiary_window_exists'] = 'This is an idio. It should be in there. If it\'s not, the tertiarty SBW section will go haywire.'
        excused_context_keys['removed_target'] = 'This is an idio. It should be in there. If it\'s not, the other removed targets section will be empty.'

        self.get_rule('PIPE-000').add_report('Email Helpers', [H2('ATS Backup Concurrence')], append=True)
        for username, member in mpst_team_client.members.items():
            if username == mpst_team_client.current_user().username: continue
            self.get_rule('PIPE-000').add_report('Email Helpers', [H3(f'If backup is {member.get_full_name()}:')], append=True)
            email_recipients = f'oco2mpst@list.jpl.nasa.gov'
            email_subject = f'Review Comments for {scr_client.filepath.name}'
            email_body = f'Hi {member.first_name}-<br /><br />'
            email_body += f'Attached are my review notes for the {week_directory_client.planning_week} ATS (Review Comments for {scr_client.filepath.name}). I need concurrence on my findings from you.  Please also review the switch buffer times.<br /><br />'
            email_body += f'Thank you,<br />{mpst_team_client.current_user().first_name}'
            self.get_rule('PIPE-000').add_report('Email Helpers', [P(email_recipients), P(email_subject), P(email_body)], append=True)

        self.get_rule('PIPE-000').add_report('Email Helpers', [H2('Approved review notes to MOM')], append=True)
        email_recipients = f'oco2@ngc.com; oco2mpst@list.jpl.nasa.gov'
        email_subject = f'MPST review of the ATS {scr_client.filepath.name}'
        email_body = f'Hello MOM,<br /><br />'
        email_body += f'A complete MPST review of the ATS {scr_client.filepath.name} (for {week_directory_client.planning_week}) has been completed and review comments are attached.'
        email_body += f'<br /><br />'
        email_body += f'Please let us know if you have any questions/comments.'
        email_body += f'<br />'
        wiki_url = f'https://wiki.jpl.nasa.gov/display/OCO2MOS/{week_directory_client.planning_week[:4]}+ATS+Reviews'
        email_body += f'MOS ATS Review Wiki:  <a href="{wiki_url}">{wiki_url}</a>'
        email_body += f'<br /><br />'
        email_body += f'Many thanks,'
        email_body += f'<br />'
        email_body += f'{mpst_team_client.current_user().first_name}'

        self.get_rule('PIPE-000').add_report('Email Helpers', [P(email_recipients), P(email_subject), P(email_body)], append=True)

        ########################################################################
        # First page summary
        #    Note that the entire first page isn't calculated here.
        #    There's also some switch buffer stuff that shows up on the
        #    first page, but that gets calculated in the switch buffer
        #    sections.
        ########################################################################

        self.add_field_to_context('scr_file_name', scr_client.filepath.name)
        self.add_field_to_context('patch_ats_directory', week_directory_client.planning_week)
        self.add_field_to_context('plan_period_start_utc', scr_client.metadata['Plan period start'])
        self.add_field_to_context('plan_period_start_mdy', scr_client.records[0]['Time'].strftime('%A %B %d, %Y'))
        self.add_field_to_context('plan_period_stop_utc', scr_client.metadata['Plan period stop'])
        self.add_field_to_context('plan_period_stop_mdy', scr_client.records[-1]['Time'].strftime('%A %B %d, %Y'))
        self.add_field_to_context('first_orbit', str(expanded_ats_client.records[0]['Orbit #']))
        self.add_field_to_context('last_orbit', str(expanded_ats_client.records[-1]['Orbit #']))
        self.add_field_to_context('car_no', str(car_client.car_no))
        self.add_field_to_context('buffer_name', car_client.buffer_name)
        url = nephthys_client.get_ats_report_link()
        self.add_hyperlink_to_context('ats_report_url', url, template)

        sbw_contacts = switch_buffer_client.unique('Contact')
        ########################################################################
        # Primary Switch Buffer
        ########################################################################
        primary_sbws = switch_buffer_client.eq('Contact', sbw_contacts[0])
        primary_switch_buffer_times = [{
            'Index': ii+1,
            'Start Time': primary_sbws['Start Time'].strftime('%H:%M:%S'),
            'Stop Time': primary_sbws['Stop Time'].strftime('%H:%M:%S')
        } for ii, primary_sbws in enumerate(primary_sbws)]

        url = nephthys_client.get_compare_sequences_page_link(primary_sbws[0]['Orbit'])
        self.add_hyperlink_to_context('prime_switch_buffer_url', url, template)
        self.add_field_to_context('primary_switch_buffer_orbit', primary_sbws[0]['Orbit'],)
        self.add_field_to_context('primary_switch_buffer_station', primary_sbws[0]['Station'],)
        self.add_list_to_context('primary_switch_buffer_times', primary_switch_buffer_times)
        self.add_field_to_context('primary_switch_buffer_duration', primary_sbws[0]['Duration (NG)'],)
        self.add_field_to_context('primary_switch_buffer_type', primary_sbws[0]['Type'],)
        self.add_field_to_context('primary_switch_buffer_aos', primary_sbws[0]['AOS (NG)'].strftime('%Y/%j %H:%M:%Sz'))
        self.add_field_to_context('primary_switch_buffer_los', primary_sbws[0]['LOS (NG)'].strftime('%Y/%j %H:%M:%Sz'))
        self.add_field_to_context('primary_switch_buffer_date', primary_sbws[0]['AOS (NG)'].strftime('%Y-%m-%d'))


        ########################################################################
        # Secondary Switch Buffer
        ########################################################################
        secondary_sbws = switch_buffer_client.eq('Contact', sbw_contacts[1])
        secondary_switch_buffer_times = [{
            'Index': ii+1,
            'Start Time': secondary_sbws['Start Time'].strftime('%H:%M:%S'),
            'Stop Time': secondary_sbws['Stop Time'].strftime('%H:%M:%S')
        } for ii, secondary_sbws in enumerate(secondary_sbws)]

        url = nephthys_client.get_compare_sequences_page_link(secondary_sbws[0]['Orbit'])
        self.add_hyperlink_to_context('secondary_switch_buffer_url', url, template)
        self.add_field_to_context('secondary_switch_buffer_orbit', secondary_sbws[0]['Orbit'],)
        self.add_field_to_context('secondary_switch_buffer_station', secondary_sbws[0]['Station'],)
        self.add_list_to_context('secondary_switch_buffer_times', secondary_switch_buffer_times)
        self.add_field_to_context('secondary_switch_buffer_duration', secondary_sbws[0]['Duration (NG)'],)
        self.add_field_to_context('secondary_switch_buffer_type', secondary_sbws[0]['Type'],)
        self.add_field_to_context('secondary_switch_buffer_aos', secondary_sbws[0]['AOS (NG)'].strftime('%Y/%j %H:%M:%Sz'))
        self.add_field_to_context('secondary_switch_buffer_los', secondary_sbws[0]['LOS (NG)'].strftime('%Y/%j %H:%M:%Sz'))
        self.add_field_to_context('secondary_switch_buffer_date', secondary_sbws[0]['AOS (NG)'].strftime('%Y-%m-%d'))

        ########################################################################
        # Tertiary Switch Buffer
        ########################################################################
        if len(sbw_contacts) == 3:
            tertiary_sbws = switch_buffer_client.eq('Contact', sbw_contacts[2])
            tertiary_switch_buffer_times = [{
                'Index': ii+1,
                'Start Time': tertiary_sbws['Start Time'].strftime('%H:%M:%S'),
                'Stop Time': tertiary_sbws['Stop Time'].strftime('%H:%M:%S')
            } for ii, tertiary_sbws in enumerate(tertiary_sbws)]
            url = nephthys_client.get_compare_sequences_page_link(tertiary_sbws[0]['Orbit'])
            self.add_hyperlink_to_context('tertiary_switch_buffer_url', url, template)
            self.add_field_to_context('tertiary_switch_buffer_orbit', tertiary_sbws[0]['Orbit'])
            self.add_field_to_context('tertiary_switch_buffer_station', tertiary_sbws[0]['Station'])
            self.add_list_to_context('tertiary_switch_buffer_times', tertiary_switch_buffer_times)
            self.add_field_to_context('tertiary_switch_buffer_duration', tertiary_sbws[0]['Duration (NG)'])
            self.add_field_to_context('tertiary_switch_buffer_type', tertiary_sbws[0]['Type'])
            self.add_field_to_context('tertiary_switch_buffer_aos', tertiary_sbws[0]['AOS (NG)'].strftime('%Y/%j %H:%M:%Sz'))
            self.add_field_to_context('tertiary_switch_buffer_los', tertiary_sbws[0]['LOS (NG)'].strftime('%Y/%j %H:%M:%Sz'))
            self.add_field_to_context('tertiary_switch_buffer_date', tertiary_sbws[0]['Stop Time'].strftime('%Y-%m-%d'))
            self.add_bool_to_context('tertiary_window_exists', True)
        else:
            self.add_bool_to_context('tertiary_window_exists', False)
            excused_context_keys['tertiary_switch_buffer_url'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_orbit'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_station'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_times'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_duration'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_type'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_aos'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_los'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_switch_buffer_date'] = 'Only 2 switch buffer contacts found'
            excused_context_keys['tertiary_window_exists'] = 'Only 2 switch buffer contacts found'

        self.add_field_to_context('abs_file_name', abs_client.filepath.name)
        self.add_field_to_context('added_nadir_cmd_orbits', ', '.join([x['Purpose'][-5:] for x in abs_client.records if 'Add Nadir CMD prior to TR Orbit ' in x['Purpose']]))
        self.add_field_to_context('added_open_door_cmd_orbits', ', '.join([x['Purpose'][-5:] for x in abs_client.records if 'Open Door post-TR contacts' in x['Purpose']]))


        ########################################################################
        # Violations page
        ########################################################################
        removed_targets = []
        orbits_with_targets_removed = [x['Orbit #'] for x in expanded_ats_client.contains('command note', 'Target removed')]
        for orbit in orbits_with_targets_removed: 
            orbit_richtext = RichText()
            orbit_richtext.add(orbit)
            target_name = RichText()
            target_name.add(odf_client.eq('Absolute Orbit #', orbit, exactly=1)['Target'])
            removed_targets.append({
                            'Orbit': orbit_richtext,
                            'Target Name': target_name,
                            })

        oef_removed_targets = []
        for target in oef_client.eq('Event', 'TargetVisibleStop').eq('Removed', True):
            removed_target_richtext = RichText()
            removed_target_richtext.add(target['orbit'])
            removed_target_richtext.add(target['targetID'])
            removed_target_richtext.add('USER TO FILL OUT REASON', color='red')
            oef_removed_targets.append(removed_target_richtext)
        if len(oef_removed_targets) == 0: 
            removed_target_richtext = RichText()
            removed_target_richtext.add('None')
            oef_removed_targets = [removed_target_richtext]

        rad_trend_records = abs_client.contains('Purpose', 'Radiometric Trend')
        if len(rad_trend_records):
            rad_trend_orbits = ' and '.join([str(o) for o in rad_trend_records.unique('Orbit')])
            #This assumes there will only ever be 1 or 2 Rad trends.
            #It will look goofy if there are ever more, but that should be rare.
            self.add_field_to_context('rad_trend_message', f'Rad trends in ABS on orbit(s) {rad_trend_orbits} are acceptable')
        else:
            self.add_field_to_context('rad_trend_message', 'No rad trends in ABS. Human to remove this line.')
        self.add_field_to_context('rad_trend_message', ' (TBC)', color='red')

        if maneuver_client.filepath is not None:
            maneuer_orbit = expanded_ats_client.eq('ACS Mode', 'THRUST')[0]['Orbit #']
            self.add_field_to_context('maneuver_message', f'{maneuver_client.filepath.name} for ')
            self.add_field_to_context('maneuver_message', f'[Maneuver Name]', color='red')
            self.add_field_to_context('maneuver_message', f'on orbit {maneuer_orbit} is ', color='red')
            self.add_field_to_context('maneuver_message', f'acceptable.', color='red')
        else:
            self.add_field_to_context('maneuver_message', f'No maneuver file present')

        if True:#maneuver_client.filepath is not None:
            self.add_field_to_context('lunar_cal_message_abs', f'Lunar cal in ABS on orbit ')
            self.add_field_to_context('lunar_cal_message_abs', f'XXXXX', color='red')
            self.add_field_to_context('lunar_cal_message_abs', f' is acceptable.')
            self.add_field_to_context('lunar_cal_message_abs', f' (TBC)', color='red')

            self.add_field_to_context('lunar_cal_message_lnr', f'Lunar cal in LNR file on orbit ')
            self.add_field_to_context('lunar_cal_message_lnr', f'XXXXX', color='red')
            self.add_field_to_context('lunar_cal_message_lnr', f' is acceptable.')
            self.add_field_to_context('lunar_cal_message_lnr', f' (TBC)', color='red')
        else:
            self.add_field_to_context('lunar_cal_message', f'Lunar cal message goes here', color='red')



        self.add_field_to_context('first_time_events', 'none')
        self.add_field_to_context('first_time_events', ' (TBC)', color='red')
        self.add_logical_check_to_context('hard_violations',
            len(violations_table.eq('Type', 'Hard')) == 0,
            'none',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('x_band_sun_pointing_violations',
            len(violations_table.eq('Type', 'X-band Hard')) == 0,
            'none',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_field_to_context('star_tracker_hard_violations', '(TBC)', color='red')
        self.add_logical_check_to_context('soft_violations',
            len(violations_table.eq('Type', 'Soft')) == 0,
            'none',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_field_to_context('n_removed_targets', len(removed_targets))
        self.add_field_to_context('target_word', 'target' if len(removed_targets) == 1 else 'targets')
        self.add_list_to_context('removed_targets', removed_targets)
        self.add_list_to_context('oef_removed_targets', oef_removed_targets)
        self.add_hyperlink_to_context('sun_angle_violations_url', nephthys_client.get_sun_violations_anchor(), template)
        self.add_field_to_context('required_initial_conditions', 'None Identified')
        self.add_field_to_context('required_initial_conditions', ' (TBC)', color='red')
        self.add_field_to_context('door_position_issues', 'None Identified')
        self.add_field_to_context('door_position_issues', ' (TBC)', color='red')
        self.add_field_to_context('ccvr_file_name', ccvr_client.filepath.name)
        self.add_field_to_context('ccvr_acceptability', 'NOT YET IMPLEMENTED', color='red')
        self.add_field_to_context('acvr_file_name', acvr_client.filepath.name)

        self.add_logical_check_to_context('acvr_acceptability', 
            sum([x['Rule ID'] not in ACCEPTABLE_ACVR_HARD_VIOLATIONS + ACCEPTABLE_ACVR_SOFT_VIOLATIONS for x in acvr_client.records]) == 0,
            'Acceptable',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('oco_ac_004_acceptability',
            sum([x['Rule ID'] in ['OCO-AC-004'] for x in acvr_client.records]) == 0,
            'Acceptable',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('oco_ac_006_acceptability',
            sum([x['Rule ID'] in ['OCO-AC-006'] for x in acvr_client.records]) == 0,
            'Acceptable',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('oco_ac_007_acceptability',
            sum([x['Rule ID'] in ['OCO-AC-007'] for x in acvr_client.records]) == 0,
            'Acceptable',
            'MANUAL REVIEW NEEDED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('oco_ac_008_acceptability_polar_cal',
            False,
            'Acceptable',
            'NOT YET IMPLEMENTED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('oco_ac_008_acceptability_end_of_xband',
            False,
            'Acceptable',
            'NOT YET IMPLEMENTED',
            expected_evaluation=True
            )
        self.add_logical_check_to_context('oco_ac_008_acceptability_targets_not_enabled',
            False,
            'Acceptable',
            'NOT YET IMPLEMENTED',
            expected_evaluation=True
            )
        self.add_field_to_context('gps_diff', int((expanded_ats_client.records[0]['Time (GPS)'] - expanded_ats_client.records[0]['Time (UTC)']).total_seconds()))
        self.add_field_to_context('reviewer_name', mpst_team_client.current_user().get_full_name())
        self.add_field_to_context('command_conference_date', datetime.now().strftime('%m/%d/%Y'))

        for var in self.context.keys():
            if var not in expected_non_record_variables and var not in expected_record_variables.keys():
                logger.warning(f'{var} is in self.context, but not in docx template')
            elif isinstance(self.context[var], list):
                if len(self.context[var]) == 0 or not isinstance(self.context[var][0], dict): continue
                if len(self.context[var]):
                    a = [x for x in self.context[var][0].keys() if x in expected_record_variables[var]]
                    b = [x for x in expected_record_variables[var] if x in self.context[var][0].keys()]
                    c = [x for x in expected_record_variables[var] if x not in self.context[var][0].keys()]
                    if len(a) != len(self.context[var][0]) or len(b) != len(expected_record_variables[var]):
                        logger.warning(f'Records in {var} are {list(self.context[var][0].keys())}. Template expects {expected_record_variables[var]}. Adding missing keys as UNKNOWN.')
                        unknown = RichText()
                        unknown.add('UNKNOWN', color='red')
                        for record in self.context[var]:
                            for key in c: 
                                record[key] = unknown
                                print(record)
                else:
                    logger.warning(f'{var} in self.context contains an empty list. Is that really what you want?')

        for var in expected_non_record_variables:
            if var not in self.context.keys():
                if var in excused_context_keys.keys():
                    logger.info(f'{var} is in docx template, but not in self.context. This is OK for the following reason: "{excused_context_keys[var]}". Adding UNKNOWN to context for this variable, but downgrading log message from WARNING to INFO')
                else:
                    logger.warning(f'{var} is in docx template, but not in self.context. Adding UNKNOWN to context for this variable')
                self.add_field_to_context(var, 'UNKNOWN', color='red')

        template.render(self.context)
        template.save(self.output_filepath)

    def add_hyperlink_to_context(self, name, url, template, display_text=None):
        if display_text is None: display_text = url
        self.add_field_to_context(name, url, url_id=template.build_url_id(url), color='blue')

    def add_bool_to_context(self, name, status):
        self.context[name] = status

    def add_subdoc_to_context(self, name, subdoc):
        self.context[name] = subdoc

    def add_field_to_context(self, name, msg, **kwargs):
        if name in self.context.keys():
            rich_text = self.context[name]
        else:
            rich_text = RichText()
        if msg is None:
            rich_text.add('UNKNOWN', **kwargs)
        else:
            rich_text.add(msg, **kwargs)

        self.context[name] = rich_text

    def add_list_to_context(self, name, lst):
        self.context[name] = lst

    def add_logical_check_to_context(self, name, logical_condition, pass_msg, fail_msg, expected_evaluation=True):
        if logical_condition == expected_evaluation:
            self.add_field_to_context(name, pass_msg)
        else:
            self.add_field_to_context(name, fail_msg, color='red')

